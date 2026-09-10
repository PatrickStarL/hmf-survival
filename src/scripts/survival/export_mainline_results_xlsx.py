#!/usr/bin/env python3
"""
Export formal mainline survival results to an XLSX workbook.
"""

from __future__ import annotations

import argparse
import json
import math
from datetime import datetime
from pathlib import Path

import pandas as pd


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export formal mainline results to XLSX.")
    parser.add_argument(
        "--run-root",
        type=Path,
        required=True,
        help="Formal run root containing reports/run_status.tsv and reports/run_config.txt.",
    )
    parser.add_argument(
        "--output-xlsx",
        type=Path,
        required=True,
        help="Path to the output XLSX workbook.",
    )
    return parser.parse_args()


def parse_run_config(path: Path) -> dict[str, str]:
    config: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or "=" not in line:
            continue
        key, value = line.split("=", 1)
        config[key] = value
    return config


def read_status_table(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, sep="\t")
    for col in ["seed", "fold"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").astype("Int64")
    return df


def read_summary_row(path_str: str) -> dict[str, object]:
    path = Path(path_str)
    if not path.is_file():
        return {
            "summary_exists": False,
            "bag_size_test": math.nan,
            "surv_loss_test": math.nan,
            "loss_test": math.nan,
            "contrastive_loss_test": math.nan,
            "c_index_test": math.nan,
        }

    df = pd.read_csv(path)
    if df.empty:
        return {"summary_exists": False}

    row = df.iloc[0].to_dict()
    row["summary_exists"] = True
    return row


def build_runs_table(status_df: pd.DataFrame) -> pd.DataFrame:
    merged_rows = []
    for row in status_df.to_dict(orient="records"):
        summary_info = read_summary_row(str(row.get("summary_csv", "")))
        merged_rows.append({**row, **summary_info})
    runs_df = pd.DataFrame(merged_rows)
    numeric_cols = [
        "bag_size_test",
        "surv_loss_test",
        "loss_test",
        "contrastive_loss_test",
        "c_index_test",
    ]
    for col in numeric_cols:
        if col in runs_df.columns:
            runs_df[col] = pd.to_numeric(runs_df[col], errors="coerce")
    return runs_df


def summarize_group(runs_df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    grouped = (
        runs_df.groupby(group_cols, dropna=False)["c_index_test"]
        .agg(
            n="count",
            mean_c_index_test="mean",
            std_c_index_test=lambda s: s.std(ddof=0),
            min_c_index_test="min",
            max_c_index_test="max",
        )
        .reset_index()
    )
    return grouped.sort_values(group_cols).reset_index(drop=True)


def build_overview(config: dict[str, str], status_df: pd.DataFrame, runs_df: pd.DataFrame) -> pd.DataFrame:
    start_raw = config.get("start_utc", "")
    end_raw = config.get("end_utc", "")
    duration_text = ""
    if start_raw and end_raw:
        start_dt = datetime.strptime(start_raw, "%Y-%m-%d %H:%M:%S UTC")
        end_dt = datetime.strptime(end_raw, "%Y-%m-%d %H:%M:%S UTC")
        duration_text = str(end_dt - start_dt)

    done_df = runs_df[runs_df["status"] == "DONE"].copy()
    by_dataset = summarize_group(done_df, ["dataset"])
    best_row = done_df.sort_values("c_index_test", ascending=False).iloc[0]
    worst_row = done_df.sort_values("c_index_test", ascending=True).iloc[0]

    overview_rows = [
        ("run_root", config.get("run_root", "")),
        ("profile", config.get("profile", "")),
        ("target_col", config.get("target_col", "")),
        ("datasets", config.get("datasets", "")),
        ("seeds", config.get("seeds", "")),
        ("folds", config.get("folds", "")),
        ("start_utc", start_raw),
        ("end_utc", end_raw),
        ("duration", duration_text),
        ("total_rows", int(len(status_df))),
        ("done", int((status_df["status"] == "DONE").sum())),
        ("skip", int((status_df["status"] == "SKIP").sum())),
        ("fail", int((status_df["status"] == "FAIL").sum())),
        ("overall_mean_c_index_test", float(done_df["c_index_test"].mean())),
        ("overall_std_c_index_test", float(done_df["c_index_test"].std(ddof=0))),
        ("best_dataset_by_mean", str(by_dataset.sort_values("mean_c_index_test", ascending=False).iloc[0]["dataset"])),
        ("best_dataset_mean_c_index_test", float(by_dataset["mean_c_index_test"].max())),
        ("worst_dataset_by_mean", str(by_dataset.sort_values("mean_c_index_test", ascending=True).iloc[0]["dataset"])),
        ("worst_dataset_mean_c_index_test", float(by_dataset["mean_c_index_test"].min())),
        ("best_run_task", str(best_row["task"])),
        ("best_run_dataset", str(best_row["dataset"])),
        ("best_run_seed", int(best_row["seed"])),
        ("best_run_fold", int(best_row["fold"])),
        ("best_run_c_index_test", float(best_row["c_index_test"])),
        ("worst_run_task", str(worst_row["task"])),
        ("worst_run_dataset", str(worst_row["dataset"])),
        ("worst_run_seed", int(worst_row["seed"])),
        ("worst_run_fold", int(worst_row["fold"])),
        ("worst_run_c_index_test", float(worst_row["c_index_test"])),
    ]
    return pd.DataFrame(overview_rows, columns=["item", "value"])


def flatten_cv_summary(path: Path) -> pd.DataFrame:
    if not path.is_file():
        return pd.DataFrame(columns=["section", "key_1", "key_2", "metric", "value"])

    data = json.loads(path.read_text(encoding="utf-8"))
    rows: list[dict[str, object]] = []

    for top_key, top_value in data.items():
        if isinstance(top_value, dict):
            for key_1, value_1 in top_value.items():
                if isinstance(value_1, dict):
                    for key_2, value_2 in value_1.items():
                        if isinstance(value_2, dict):
                            for metric, metric_value in value_2.items():
                                rows.append(
                                    {
                                        "section": top_key,
                                        "key_1": key_1,
                                        "key_2": key_2,
                                        "metric": metric,
                                        "value": metric_value,
                                    }
                                )
                        else:
                            rows.append(
                                {
                                    "section": top_key,
                                    "key_1": key_1,
                                    "key_2": "",
                                    "metric": key_2,
                                    "value": value_2,
                                }
                            )
                else:
                    rows.append(
                        {
                            "section": top_key,
                            "key_1": "",
                            "key_2": "",
                            "metric": key_1,
                            "value": value_1,
                        }
                    )
        else:
            rows.append(
                {
                    "section": top_key,
                    "key_1": "",
                    "key_2": "",
                    "metric": "",
                    "value": top_value,
                }
            )
    return pd.DataFrame(rows)


def autofit_columns(output_xlsx: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(output_xlsx)
    for ws in wb.worksheets:
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                if len(value) > max_len:
                    max_len = len(value)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 60)
        ws.freeze_panes = "A2"
    wb.save(output_xlsx)


def main() -> None:
    args = parse_args()
    reports_dir = args.run_root / "reports"
    status_path = reports_dir / "run_status.tsv"
    config_path = reports_dir / "run_config.txt"
    cv_summary_path = reports_dir / "cv_summary_mainline_twostage_5seed5fold.json"

    config = parse_run_config(config_path)
    status_df = read_status_table(status_path)
    runs_df = build_runs_table(status_df)
    done_df = runs_df[runs_df["status"] == "DONE"].copy().sort_values(["dataset", "seed", "fold"]).reset_index(drop=True)

    overview_df = build_overview(config, status_df, done_df)
    config_df = pd.DataFrame(
        [{"key": key, "value": value} for key, value in config.items()]
    )
    by_dataset_df = summarize_group(done_df, ["dataset"])
    by_dataset_seed_df = summarize_group(done_df, ["dataset", "seed"])
    by_dataset_fold_df = summarize_group(done_df, ["dataset", "fold"])
    top_runs_df = done_df.sort_values("c_index_test", ascending=False).head(20).reset_index(drop=True)
    bottom_runs_df = done_df.sort_values("c_index_test", ascending=True).head(20).reset_index(drop=True)
    cv_summary_flat_df = flatten_cv_summary(cv_summary_path)

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(args.output_xlsx, engine="openpyxl") as writer:
        overview_df.to_excel(writer, sheet_name="overview", index=False)
        config_df.to_excel(writer, sheet_name="run_config", index=False)
        status_df.to_excel(writer, sheet_name="run_status", index=False)
        done_df.to_excel(writer, sheet_name="runs_done", index=False)
        by_dataset_df.to_excel(writer, sheet_name="by_dataset", index=False)
        by_dataset_seed_df.to_excel(writer, sheet_name="by_dataset_seed", index=False)
        by_dataset_fold_df.to_excel(writer, sheet_name="by_dataset_fold", index=False)
        top_runs_df.to_excel(writer, sheet_name="top_20_runs", index=False)
        bottom_runs_df.to_excel(writer, sheet_name="bottom_20_runs", index=False)
        cv_summary_flat_df.to_excel(writer, sheet_name="cv_summary_flat", index=False)

    autofit_columns(args.output_xlsx)
    print(args.output_xlsx)


if __name__ == "__main__":
    main()
