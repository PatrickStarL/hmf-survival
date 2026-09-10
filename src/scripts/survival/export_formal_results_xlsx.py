#!/usr/bin/env python3
"""
Export current formal experiment results into a single XLSX workbook.

Included experiments:
- Mainline 5-cancer formal result
- P1 fusion structure comparison (latest unique rows after rerun)
- P1 modality contribution ablation (current snapshot; may be in progress)
"""

from __future__ import annotations

import argparse
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

import pandas as pd


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export formal experiment results to one XLSX workbook.")
    parser.add_argument(
        "--mainline-run-root",
        type=Path,
        default=None,
        help="Mainline run root. Defaults to latest 5-cancer mainline root.",
    )
    parser.add_argument(
        "--structure-run-root",
        type=Path,
        default=None,
        help="Structure compare run root. Defaults to latest formal structure compare root.",
    )
    parser.add_argument(
        "--modality-run-root",
        type=Path,
        default=None,
        help="Modality ablation run root. Defaults to latest formal modality-ablation root.",
    )
    parser.add_argument(
        "--output-xlsx",
        type=Path,
        default=None,
        help="Output workbook path. Defaults to /path/to/your/data_root/formal_results_snapshot_<ts>.xlsx",
    )
    return parser.parse_args()


def latest_matching(glob_pattern: str) -> Path | None:
    matches = sorted(Path("/").glob(glob_pattern.lstrip("/")), key=lambda p: p.stat().st_mtime, reverse=True)
    return matches[0] if matches else None


def default_output_path() -> Path:
    ts = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    return Path(f"/path/to/your/data_root/formal_results_snapshot_{ts}.xlsx")


def resolve_run_root(given: Path | None, glob_pattern: str, label: str) -> Path | None:
    if given is not None:
        return given
    found = latest_matching(glob_pattern)
    if found is None:
        print(f"[warn] no run root found for {label} via pattern: {glob_pattern}")
    return found


def parse_run_config(path: Path) -> dict[str, str]:
    if not path.is_file():
        return {}
    config: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or "=" not in line:
            continue
        key, value = line.split("=", 1)
        config[key] = value
    return config


def read_status_table(path: Path) -> pd.DataFrame:
    if not path.is_file():
        return pd.DataFrame()
    df = pd.read_csv(path, sep="\t")
    if "timestamp_utc" in df.columns:
        df["timestamp_utc_dt"] = pd.to_datetime(df["timestamp_utc"], errors="coerce", utc=True)
    for col in ["seed", "fold", "c_index_test"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def latest_unique_rows(df: pd.DataFrame, key_cols: list[str]) -> pd.DataFrame:
    if df.empty:
        return df.copy()
    work = df.copy()
    if "timestamp_utc_dt" in work.columns:
        work = work.sort_values(["timestamp_utc_dt"] + key_cols, kind="stable")
    else:
        work = work.reset_index(drop=False).sort_values(["index"] + key_cols, kind="stable")
    latest = work.groupby(key_cols, dropna=False, as_index=False).tail(1)
    return latest.reset_index(drop=True)


def summarize_group(df: pd.DataFrame, group_cols: list[str], value_col: str = "c_index_test") -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=group_cols + ["n", "mean_c_index_test", "std_c_index_test", "min_c_index_test", "max_c_index_test"])
    done = df[df["status"] == "DONE"].copy()
    if done.empty:
        return pd.DataFrame(columns=group_cols + ["n", "mean_c_index_test", "std_c_index_test", "min_c_index_test", "max_c_index_test"])
    out = (
        done.groupby(group_cols, dropna=False)[value_col]
        .agg(
            n="count",
            mean_c_index_test="mean",
            std_c_index_test=lambda s: s.std(ddof=0),
            min_c_index_test="min",
            max_c_index_test="max",
        )
        .reset_index()
        .sort_values(group_cols)
        .reset_index(drop=True)
    )
    return out


def build_experiment_registry(rows: Iterable[dict[str, object]]) -> pd.DataFrame:
    return pd.DataFrame(list(rows), columns=[
        "experiment",
        "run_root",
        "status_tsv",
        "run_config",
        "latest_total",
        "latest_done",
        "latest_fail",
        "latest_skip",
        "note",
    ])


def experiment_summary_row(
    experiment: str,
    run_root: Path | None,
    status_df: pd.DataFrame,
    latest_df: pd.DataFrame,
    run_config_path: Path | None,
    note: str,
) -> dict[str, object]:
    return {
        "experiment": experiment,
        "run_root": "" if run_root is None else str(run_root),
        "status_tsv": "" if run_root is None else str(run_root / "reports" / "run_status.tsv"),
        "run_config": "" if run_config_path is None else str(run_config_path),
        "latest_total": int(len(latest_df)),
        "latest_done": int((latest_df["status"] == "DONE").sum()) if not latest_df.empty else 0,
        "latest_fail": int((latest_df["status"] == "FAIL").sum()) if not latest_df.empty else 0,
        "latest_skip": int((latest_df["status"] == "SKIP").sum()) if not latest_df.empty else 0,
        "note": note,
    }


def flatten_config(config: dict[str, str]) -> pd.DataFrame:
    return pd.DataFrame([{"key": k, "value": v} for k, v in config.items()])


def align_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=columns)
    work = df.copy()
    for col in columns:
        if col not in work.columns:
            work[col] = pd.NA
    return work[columns]


def add_experiment_column(df: pd.DataFrame, experiment: str) -> pd.DataFrame:
    if df.empty:
        return df.copy()
    work = df.copy()
    work.insert(0, "experiment", experiment)
    return work


def rerun_target_tables(structure_latest_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if structure_latest_df.empty:
        empty = pd.DataFrame(columns=["dataset", "seed", "fold", "status", "c_index_test", "task", "summary_csv", "log_file", "note"])
        return empty, pd.DataFrame(columns=["dataset", "n", "mean_c_index_test", "std_c_index_test", "min_c_index_test", "max_c_index_test"])

    target = structure_latest_df[
        (structure_latest_df["structure"] == "shallow_cls")
        & (structure_latest_df["dataset"].isin(["LUAD", "STAD"]))
    ].copy()
    target = target.sort_values(["dataset", "seed", "fold"]).reset_index(drop=True)
    summary = summarize_group(target, ["dataset"])
    return target, summary


def autofit_columns(output_xlsx: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(output_xlsx)
    for ws in wb.worksheets:
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                if len(value) > max_len:
                    max_len = len(value)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 60)
        ws.freeze_panes = "A2"
    wb.save(output_xlsx)


def excel_safe(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df.copy()
    drop_cols = [col for col in df.columns if col.endswith("_dt")]
    return df.drop(columns=drop_cols, errors="ignore")


def main() -> None:
    args = parse_args()

    mainline_root = resolve_run_root(
        args.mainline_run_root,
        "/path/to/your/data_root/mainline_twostage_panther_mstar_conch_5cancer_*",
        "mainline",
    )
    structure_root = resolve_run_root(
        args.structure_run_root,
        "/path/to/your/data_root/formal_p1_fusion_structure_compare_3cancer_3seed5fold_*",
        "structure_compare",
    )
    modality_root = resolve_run_root(
        args.modality_run_root,
        "/path/to/your/data_root/formal_p1_modality_contribution_ablation_3cancer_3seed5fold_*",
        "modality_ablation",
    )
    output_xlsx = args.output_xlsx or default_output_path()

    registry_rows: list[dict[str, object]] = []

    mainline_status = read_status_table(mainline_root / "reports" / "run_status.tsv") if mainline_root else pd.DataFrame()
    mainline_latest = latest_unique_rows(mainline_status, ["dataset", "seed", "fold"]) if not mainline_status.empty else pd.DataFrame()
    mainline_config_path = (mainline_root / "reports" / "run_config.txt") if mainline_root else None
    mainline_config = parse_run_config(mainline_config_path) if mainline_config_path else {}
    registry_rows.append(
        experiment_summary_row(
            "mainline_5cancer",
            mainline_root,
            mainline_status,
            mainline_latest,
            mainline_config_path,
            "official 5-cancer mainline",
        )
    )

    structure_status = read_status_table(structure_root / "reports" / "run_status.tsv") if structure_root else pd.DataFrame()
    structure_latest = latest_unique_rows(structure_status, ["structure", "dataset", "seed", "fold"]) if not structure_status.empty else pd.DataFrame()
    structure_config_path = (structure_root / "reports" / "run_config.txt") if structure_root else None
    structure_config = parse_run_config(structure_config_path) if structure_config_path else {}
    registry_rows.append(
        experiment_summary_row(
            "p1_fusion_structure_compare",
            structure_root,
            structure_status,
            structure_latest,
            structure_config_path,
            "latest unique rows reflect rerun-backfilled final results",
        )
    )

    modality_status = read_status_table(modality_root / "reports" / "run_status.tsv") if modality_root else pd.DataFrame()
    modality_latest = latest_unique_rows(modality_status, ["ablation", "dataset", "seed", "fold"]) if not modality_status.empty else pd.DataFrame()
    modality_config_path = (modality_root / "reports" / "run_config.txt") if modality_root else None
    modality_config = parse_run_config(modality_config_path) if modality_config_path else {}
    registry_rows.append(
        experiment_summary_row(
            "p1_modality_ablation",
            modality_root,
            modality_status,
            modality_latest,
            modality_config_path,
            "current snapshot; experiment still in progress if latest_done < latest_total",
        )
    )

    registry_df = build_experiment_registry(registry_rows)

    mainline_by_dataset = summarize_group(mainline_latest, ["dataset"])
    mainline_by_dataset_seed = summarize_group(mainline_latest, ["dataset", "seed"])
    mainline_by_dataset_fold = summarize_group(mainline_latest, ["dataset", "fold"])

    structure_by_structure = summarize_group(structure_latest, ["structure"])
    structure_by_structure_dataset = summarize_group(structure_latest, ["structure", "dataset"])
    rerun_target_rows, rerun_target_summary = rerun_target_tables(structure_latest)

    modality_by_ablation = summarize_group(modality_latest, ["ablation"])
    modality_by_ablation_dataset = summarize_group(modality_latest, ["ablation", "dataset"])

    combined_cols = [
        "experiment",
        "timestamp_utc",
        "status",
        "structure",
        "ablation",
        "dataset",
        "seed",
        "fold",
        "task",
        "c_index_test",
        "summary_csv",
        "log_file",
        "note",
    ]
    combined_latest = pd.concat(
        [
            align_columns(add_experiment_column(mainline_latest, "mainline_5cancer"), combined_cols),
            align_columns(add_experiment_column(structure_latest, "p1_fusion_structure_compare"), combined_cols),
            align_columns(add_experiment_column(modality_latest, "p1_modality_ablation"), combined_cols),
        ],
        ignore_index=True,
    )

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        excel_safe(registry_df).to_excel(writer, sheet_name="overview", index=False)

        if not mainline_status.empty:
            excel_safe(flatten_config(mainline_config)).to_excel(writer, sheet_name="mainline_config", index=False)
            excel_safe(mainline_status).to_excel(writer, sheet_name="mainline_status_raw", index=False)
            excel_safe(mainline_latest.sort_values(["dataset", "seed", "fold"])).to_excel(writer, sheet_name="mainline_latest", index=False)
            excel_safe(mainline_by_dataset).to_excel(writer, sheet_name="mainline_by_dataset", index=False)
            excel_safe(mainline_by_dataset_seed).to_excel(writer, sheet_name="mainline_by_ds_seed", index=False)
            excel_safe(mainline_by_dataset_fold).to_excel(writer, sheet_name="mainline_by_ds_fold", index=False)

        if not structure_status.empty:
            excel_safe(flatten_config(structure_config)).to_excel(writer, sheet_name="structure_config", index=False)
            excel_safe(structure_status).to_excel(writer, sheet_name="structure_status_raw", index=False)
            excel_safe(structure_latest.sort_values(["structure", "dataset", "seed", "fold"])).to_excel(writer, sheet_name="structure_latest", index=False)
            excel_safe(structure_by_structure).to_excel(writer, sheet_name="structure_by_type", index=False)
            excel_safe(structure_by_structure_dataset).to_excel(writer, sheet_name="structure_by_type_ds", index=False)
            excel_safe(rerun_target_rows).to_excel(writer, sheet_name="rerun_targets_rows", index=False)
            excel_safe(rerun_target_summary).to_excel(writer, sheet_name="rerun_targets_sum", index=False)

        if not modality_status.empty:
            excel_safe(flatten_config(modality_config)).to_excel(writer, sheet_name="modality_config", index=False)
            excel_safe(modality_status).to_excel(writer, sheet_name="modality_status_raw", index=False)
            excel_safe(modality_latest.sort_values(["ablation", "dataset", "seed", "fold"])).to_excel(writer, sheet_name="modality_latest", index=False)
            excel_safe(modality_by_ablation).to_excel(writer, sheet_name="modality_by_ablt", index=False)
            excel_safe(modality_by_ablation_dataset).to_excel(writer, sheet_name="modality_by_ablt_ds", index=False)

        excel_safe(combined_latest).to_excel(writer, sheet_name="all_latest", index=False)

    autofit_columns(output_xlsx)
    print(output_xlsx)


if __name__ == "__main__":
    main()
